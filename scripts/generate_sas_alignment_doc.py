#!/usr/bin/env python3
"""
Generates the ComplianceIQ x AWS Security Assurance Services (SAS) alignment
and roadmap Word document. Saved to ~/Documents.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------------------
# Palette (aligned to ComplianceIQ dark-slate / emerald / cyan brand)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
EMERALD = RGBColor(0x0F, 0x9D, 0x58)
TEAL = RGBColor(0x0D, 0x94, 0x88)
CYAN = RGBColor(0x06, 0x83, 0x9B)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GREY = RGBColor(0x50, 0x5A, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# ---------------------------------------------------------------------------
# Base styles
# ---------------------------------------------------------------------------
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0x20, 0x25, 0x30)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def _set_cell_text(cell, text, bold=False, color=None, size=9.5, align=None, white=False):
    cell.text = ''
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = 'Calibri'
    if white:
        run.font.color.rgb = WHITE
    elif color is not None:
        run.font.color.rgb = color

def heading(text, level=1, color=NAVY, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.bold = True
    run.font.name = 'Calibri'
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = TEAL
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), '0D9488')
        pbdr.append(bottom)
        pPr.append(pbdr)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = NAVY
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = CYAN
    return p

def body(text, bold=False, italic=False, color=None, size=10.5, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.bold = bold
    run.font.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return p

def bullet(text, bold_lead=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3 + level * 0.3)
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.font.bold = True
        r.font.size = Pt(10.5)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)
    else:
        r = p.add_run(text)
        r.font.size = Pt(10.5)
    return p

def make_table(headers, rows, col_widths=None, header_fill='0F172A'):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _shade(hdr[i], header_fill)
        _set_cell_text(hdr[i], h, bold=True, white=True, size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _set_cell_text(cells[i], val, size=9.5)
            if i == 0:
                # first column subtle emphasis
                for para in cells[i].paragraphs:
                    for run in para.runs:
                        run.font.bold = True
                        run.font.color.rgb = NAVY
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    return table

def spacer(pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    return p

# ===========================================================================
# COVER
# ===========================================================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run('ComplianceIQ')
r.font.size = Pt(30)
r.font.bold = True
r.font.color.rgb = NAVY
r2 = title.add_run('  |  Regulatory Observatory')
r2.font.size = Pt(18)
r2.font.color.rgb = TEAL

sub = doc.add_paragraph()
r = sub.add_run('Aligning an AI-Powered Compliance Intelligence Tool with the '
                'AWS Security Assurance Services (SAS) Engagement Lifecycle')
r.font.size = Pt(13)
r.font.italic = True
r.font.color.rgb = SLATE

meta = doc.add_paragraph()
meta.paragraph_format.space_before = Pt(10)
for label, val in [
    ('Prepared for', 'SAS Team & Broader Enterprise Support / Global Services Organization'),
    ('Purpose', 'Team objective alignment, productivity & client-value narrative, 3-horizon roadmap'),
    ('Current scope', 'MENAT region (24 sovereign jurisdictions) — architected to expand to other regions'),
    ('Document status', 'Internal working draft for team review'),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    rl = p.add_run(f'{label}:  ')
    rl.font.bold = True
    rl.font.size = Pt(10)
    rl.font.color.rgb = TEAL
    rv = p.add_run(val)
    rv.font.size = Pt(10)
    rv.font.color.rgb = SLATE

# thin rule
rule = doc.add_paragraph()
pPr = rule._p.get_or_add_pPr()
pbdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12')
bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '0D9488')
pbdr.append(bottom); pPr.append(pbdr)

# ===========================================================================
# 1. EXECUTIVE SUMMARY
# ===========================================================================
heading('1. Executive Summary', level=1)
body('ComplianceIQ is an AI-powered regulatory intelligence application that consolidates '
     'sovereign regulations, mandatory cybersecurity frameworks, data-protection decrees, and '
     'AI-governance baselines into a single, queryable knowledge surface. It currently covers the '
     '24 sovereign jurisdictions of the Middle East, North Africa and Türkiye (MENAT) region and '
     'maps national controls to global frameworks — NIST CSF 2.0, ISO/IEC 27001:2022 and CSA CCM v4.')
body('The application runs entirely inside the AWS boundary. Its AI reasoning is powered by Amazon '
     'Bedrock, and no client or regulatory data leaves AWS during model inference. This makes it a '
     'natural, policy-safe productivity accelerator for the AWS Security Assurance Services (SAS) team.')
body('This document maps ComplianceIQ directly onto the five SAS delivery pillars, articulates the '
     'direct and indirect benefits as an internal productivity tool and as a potential source of '
     'client value, and proposes a three-horizon roadmap. The central thesis is simple:', space_after=3)
body('ComplianceIQ compresses the research- and evidence-heavy phases of a compliance engagement — '
     'the parts that today consume senior consultant hours on manual regulation lookup, control '
     'mapping, and gap analysis — into minutes, while keeping every step grounded, auditable, and '
     'inside AWS.', bold=True, color=NAVY)

# ===========================================================================
# 2. HOW COMPLIANCEIQ WORKS TODAY
# ===========================================================================
heading('2. What ComplianceIQ Does Today', level=1)
body('The following capabilities are already built and operational in the application. They form the '
     'foundation for the SAS lifecycle mapping in Section 4.')

make_table(
    ['Capability', 'What it does', 'SAS relevance'],
    [
        ['Sovereign Regulation Registry',
         'Indexed statutory instruments across 24 MENAT jurisdictions with authority, status, effective dates and penalties.',
         'Regulatory mapping & scoping'],
        ['Controls Crosswalk',
         'Bi-directional mapping of national controls to NIST CSF 2.0, ISO/IEC 27001:2022 and CSA CCM v4.',
         'Framework selection & gap assessment'],
        ['AI Control Clause Interpreter',
         'Breaks a control into People / Process / Technical actions plus an auditor evidence checklist.',
         'Audit readiness & remediation'],
        ['AI Policy Redlining',
         'Scores a client policy against a chosen regulation, flags missing mandatory clauses, drafts fixes.',
         'Gap assessment & remediation'],
        ['Maturity Heatmap & Gap Memos',
         'Country / sector maturity scoring with AI-generated executive gap-analysis memos.',
         'Advisory & scoping'],
        ['Regulatory Radar & Live News',
         'Tracks upcoming drafts, amendments and enforcement changes with web-grounded intelligence.',
         'Continuous compliance / drift'],
        ['Requirement Confidence Engine',
         'Classifies each requirement as Mandatory / Guideline / Conditional with a confidence score.',
         'Scoping & prioritisation'],
        ['Link-Integrity & Source Audit',
         'Automated reachability checks on official regulatory source URLs and gazette PDFs.',
         'Evidence provenance & assurance'],
        ['Audit-Ready Export',
         'One-click PDF / CSV / JSON dossiers of regulations, controls and crosswalks.',
         'Evidence packaging'],
        ['RBAC & Audit Trail',
         'Role separation (analyst / manager / auditor / admin) with an immutable action ledger.',
         'Governance & separation of duties'],
    ],
    col_widths=[1.7, 3.4, 1.7],
)

body('Technical note: the AI layer is served by Amazon Bedrock (Amazon Nova family) within the '
     'customer AWS account; web-grounding is optional and pluggable. Because inference stays inside '
     'AWS, the tool is compatible with the data-handling posture SAS engagements require.',
     italic=True, size=9.5, color=GREY, space_after=8)

# ===========================================================================
# 3. THE SAS ENGAGEMENT LIFECYCLE
# ===========================================================================
heading('3. The SAS Engagement Lifecycle (Reference)', level=1)
body('AWS Security Assurance Services delivers audit and compliance engineering across the full cloud '
     'compliance journey, combining experienced auditors with AWS technical depth. The publicly stated '
     'value themes include reducing compliance cost through automation, shortening certification '
     'timelines so compliance becomes a business enabler, embedding compliance into DevSecOps, and '
     'producing compliance playbooks that double as communication tools for auditors and regulators. '
     '(Source: aws.amazon.com/security-assurance-services — content paraphrased.)')
body('SAS delivery is organised around five pillars:', bold=True, color=NAVY, space_after=3)
for lead, rest in [
    ('Compliance Advisory — ', 'gap assessments, scoping, framework selection, and regulatory mapping.'),
    ('Compliance Engineering — ', 'automated controls, Config rules, Security Hub integration, and policy as code.'),
    ('Remediation Support — ', 'hands-on remediation of identified gaps and architecture optimisation.'),
    ('Audit Readiness — ', 'evidence packaging, playbook delivery, and support for auditor conversations.'),
    ('Continuous Compliance — ', 'monitoring, drift detection, and ongoing program support.'),
]:
    bullet(rest, bold_lead=lead)
body('Note on sourcing: several supporting references provided (internal AWS wiki, the SAS Target '
     'Operating Model on SharePoint, Highspot, and Broadcast) require Amazon single-sign-on and could '
     'not be retrieved by the document generator. The lifecycle above reflects the five deliverables '
     'supplied directly plus the public SAS site. Anyone with SSO access should validate the mapping '
     'in Section 4 against the current TOM before external circulation.',
     italic=True, size=9, color=AMBER, space_after=8)

# ===========================================================================
# 4. LIFECYCLE MAPPING
# ===========================================================================
heading('4. Mapping ComplianceIQ to the SAS Lifecycle', level=1)
body('The table below connects each SAS pillar to what ComplianceIQ can do today, and where it points '
     'next. This is the core of the alignment story.')

make_table(
    ['SAS Pillar', 'How ComplianceIQ supports it today', 'Where it goes next'],
    [
        ['Compliance Advisory',
         'Regulation registry, requirement-confidence classification, maturity heatmaps and AI gap-analysis memos let a consultant scope an engagement and pick the right framework in minutes rather than days.',
         'Client-specific scoping wizard that ingests a target profile and returns a tailored regulatory scope.'],
        ['Compliance Engineering',
         'Controls Crosswalk maps sovereign requirements to NIST / ISO / CSA, giving engineers the control set to implement; the clause interpreter translates each control into technical actions.',
         'Generate AWS Config rule / Security Hub control suggestions and policy-as-code snippets per mapped control.'],
        ['Remediation Support',
         'AI Policy Redlining pinpoints non-compliant clauses and drafts remediation language; the interpreter lists the exact People/Process/Technical fixes.',
         'Track remediation items to closure with owner, status and re-test evidence inside the app.'],
        ['Audit Readiness',
         'Auditor evidence checklists per control, provenance-checked source links, and one-click PDF/CSV/JSON dossiers package evidence and support auditor conversations.',
         'Auto-assembled, framework-specific audit playbooks and an auditor-facing read-only evidence view.'],
        ['Continuous Compliance',
         'Regulatory Radar, live grounded news and the weekly source-audit daemon surface drafts, amendments and broken/changed sources — early drift signals.',
         'Scheduled re-assessment, delta alerts against a saved baseline, and program dashboards per client.'],
    ],
    col_widths=[1.4, 3.7, 1.7],
)

# ===========================================================================
# 5. BENEFITS
# ===========================================================================
heading('5. Benefits: Productivity Tool and Client Value', level=1)

heading('5.1 Direct Benefits (Internal Productivity)', level=2, color=NAVY)
for lead, rest in [
    ('Faster scoping & research — ', 'regulation lookup, control mapping and first-pass gap analysis collapse from days of manual work to minutes, freeing senior consultants for judgement-heavy work.'),
    ('Consistency & quality — ', 'every engagement starts from the same grounded control mappings and confidence scoring, reducing variance between consultants and rework.'),
    ('Lower onboarding cost — ', 'newer team members become productive faster because the tool encodes framework knowledge and crosswalks.'),
    ('Reusable, auditable output — ', 'exports, evidence checklists and the action ledger create artefacts that can be reused across engagements and stand up to review.'),
    ('Policy-safe by design — ', 'all AI inference runs inside AWS on Amazon Bedrock, so the tool can be used on sensitive material without data leaving the AWS boundary.'),
    ('Scales the team\'s reach — ', 'the same headcount can cover more jurisdictions and more clients because the repetitive research load is automated.'),
]:
    bullet(rest, bold_lead=lead)

heading('5.2 Indirect Benefits', level=2, color=NAVY)
for lead, rest in [
    ('Knowledge retention — ', 'institutional regulatory knowledge is captured in the app rather than living only in individual consultants\' heads.'),
    ('Shorter time-to-value for clients — ', 'faster scoping and gap analysis compresses certification timelines, reinforcing the SAS message that compliance is a business enabler.'),
    ('Stronger auditor conversations — ', 'grounded citations and provenance-checked sources raise the credibility of evidence packages.'),
    ('Cross-sell signal — ', 'maturity heatmaps and radar naturally surface adjacent gaps a client may want help closing.'),
    ('Brand & thought-leadership — ', 'a polished, AWS-native compliance intelligence surface is a compelling artefact in client conversations and demos.'),
]:
    bullet(rest, bold_lead=lead)

heading('5.3 Potential Client-Facing Value', level=2, color=NAVY)
body('Over time, and subject to the appropriate legal, data-handling and go-to-market review, elements '
     'of ComplianceIQ could be exposed to clients directly:')
for lead, rest in [
    ('Self-service regulatory scoping — ', 'clients explore their applicable regulations and maturity before or between engagements.'),
    ('Living compliance dashboard — ', 'a per-client view of scope, gaps, remediation status and drift, refreshed continuously.'),
    ('Evidence & playbook portal — ', 'clients retrieve audit-ready dossiers and framework playbooks on demand.'),
    ('Productised service offering — ', 'a subscription intelligence layer that complements, rather than replaces, hands-on SAS delivery.'),
]:
    bullet(rest, bold_lead=lead)
body('Important: any client-facing exposure must preserve the SAS principle that the customer remains '
     'solely responsible for their own compliance. ComplianceIQ provides tooling, mapping and guidance '
     '— it does not determine or certify compliance on a client\'s behalf.',
     italic=True, size=9.5, color=AMBER, space_after=8)

# ===========================================================================
# 6. THREE-HORIZON ROADMAP
# ===========================================================================
heading('6. Three-Horizon Roadmap', level=1)
body('The roadmap is framed with the classic three-horizon model. Horizon 1 is live today; Horizon 2 is '
     'the immediate build priority that deepens SAS lifecycle fit; Horizon 3 is the strategic direction, '
     'including regional expansion and a potential client-facing service.')

heading('Horizon 1 — Current Capabilities (Live Today)', level=2, color=EMERALD)
for t in [
    'Regulation registry & controls crosswalk across 24 MENAT jurisdictions (NIST / ISO / CSA).',
    'AI clause interpreter, policy redlining, maturity heatmaps and gap-analysis memos on Amazon Bedrock.',
    'Regulatory radar, web-grounded news, weekly source-integrity audit.',
    'RBAC, audit trail, and audit-ready PDF / CSV / JSON export.',
]:
    bullet(t)

heading('Horizon 2 — Immediate Priority (Build Next, ~1–2 Quarters)', level=2, color=TEAL)
body('These features convert ComplianceIQ from a research accelerator into an engagement companion that '
     'tracks work through the SAS lifecycle.', size=10, color=GREY)
for lead, rest in [
    ('Engagement workspaces — ', 'a per-client container that holds scope, mapped controls, gaps, remediation items and evidence in one place.'),
    ('Remediation tracking — ', 'convert redline findings and gaps into owned, status-tracked items with re-test evidence.'),
    ('Config / Security Hub suggestions — ', 'for each mapped control, propose the relevant AWS Config rule or Security Hub standard and a policy-as-code starting point (Compliance Engineering fit).'),
    ('Auto-assembled audit playbooks — ', 'generate framework-specific evidence packs and auditor-conversation guides from a completed assessment.'),
    ('Baseline & drift alerts — ', 'save an assessment baseline and alert on regulatory or posture drift against it.'),
    ('SSO & enterprise hardening — ', 'Amazon federated login, least-privilege IAM, and CloudTrail-backed audit logging for internal rollout.'),
]:
    bullet(rest, bold_lead=lead)

heading('Horizon 3 — Future Priority (Plan Soon, Strategic)', level=2, color=CYAN)
for lead, rest in [
    ('Multi-region expansion — ', 'extend beyond MENAT to additional regions (EU, APAC, Americas) reusing the same crosswalk architecture; the data model is already region-agnostic.'),
    ('Deeper framework coverage — ', 'add PCI DSS, SOC 2, HIPAA, FedRAMP, DORA and other SAS-supported frameworks to broaden engagement fit.'),
    ('Client-facing service offering — ', 'a governed, subscription intelligence layer delivering self-service scoping, a living compliance dashboard and an evidence portal — complementing hands-on SAS delivery.'),
    ('Live AWS posture integration — ', 'read from a client\'s AWS Config / Security Hub to compare actual posture against mapped requirements for real continuous compliance.'),
    ('Agentic remediation assist — ', 'guided, human-in-the-loop generation of remediation IaC and policy-as-code, kept inside AWS.'),
    ('Benchmarking & analytics — ', 'anonymised, opt-in cross-engagement maturity benchmarking as a thought-leadership asset.'),
]:
    bullet(rest, bold_lead=lead)

# Horizon summary table
spacer()
make_table(
    ['Horizon', 'Theme', 'Primary beneficiary', 'Headline outcome'],
    [
        ['H1 — Now', 'Research & mapping acceleration', 'Internal (consultants)', 'Days of scoping/research -> minutes'],
        ['H2 — Next', 'Engagement lifecycle companion', 'Internal + early client value', 'Track scope-to-evidence in one place'],
        ['H3 — Future', 'Regional + productised service', 'Internal + direct client offering', 'New coverage & a scalable service layer'],
    ],
    col_widths=[1.1, 2.2, 1.9, 2.3],
)

# ===========================================================================
# 7. DEPENDENCIES & INTEGRATION REQUIREMENTS
# ===========================================================================
heading('7. Dependencies & Integration Requirements', level=1)
body('Each horizon carries dependencies that must be satisfied before or during delivery. These fall '
     'into four categories: (a) technical integrations, (b) client-side support, (c) artefacts/inputs '
     'needed, and (d) other non-invasive assistance. Two hard constraints apply to every integration '
     'in this roadmap and are non-negotiable:')
for lead, rest in [
    ('Non-production only — ',
     'all integrations connect exclusively to non-production / sandbox / staging AWS environments. '
     'No connection is made to a client production account.'),
    ('Read-only access only — ',
     'where the tool reads from AWS services (e.g. Security Hub, AWS Config, Conformance Packs) or '
     'third-party GRC tools (e.g. AuditBoard, if/when available), it uses least-privilege, read-only '
     'credentials. ComplianceIQ never writes, modifies, or remediates directly in a client account.'),
]:
    bullet(rest, bold_lead=lead)
body('This "observe, never touch" posture keeps ComplianceIQ non-invasive: it ingests posture and '
     'evidence signals to inform consultants, while all changes remain in the hands of the client or '
     'the SAS engineer through their own change-controlled process.',
     italic=True, size=9.5, color=GREY, space_after=8)

# ---- 7.1 Horizon 1 dependencies ----
heading('7.1 Horizon 1 Dependencies (Live Today)', level=2, color=EMERALD)
body('Horizon 1 is self-contained — it runs on curated regulatory data and AI reasoning with no client '
     'environment access required. Dependencies are minimal.', size=10, color=GREY)
make_table(
    ['Dependency type', 'Detail', 'Constraint / notes'],
    [
        ['Integrations',
         'Amazon Bedrock (AI inference) in a team-owned non-prod AWS account; optional web-grounding search API.',
         'Inference stays inside AWS; Bedrock invoked read-only via InvokeModel. No client account touched.'],
        ['Client support',
         'None required for internal use. For a client demo, only permission to discuss their in-scope frameworks.',
         'Zero access to client systems at this horizon.'],
        ['Artefacts needed',
         'Curated regulatory dataset (already built); optionally a client policy document for the redlining demo.',
         'Client documents handled per data-handling policy; can use redacted/sample docs.'],
        ['Other assistance',
         'Team-owned non-prod AWS account with Bedrock model access enabled.',
         'One-time setup; no ongoing client dependency.'],
    ],
    col_widths=[1.4, 3.6, 1.8],
)

# ---- 7.2 Horizon 2 dependencies ----
heading('7.2 Horizon 2 Dependencies (Immediate Priority)', level=2, color=TEAL)
body('Horizon 2 introduces the first read-only connections to a client (or SAS-managed) non-production '
     'AWS environment, plus richer engagement artefacts. This is where integration detail matters most.',
     size=10, color=GREY)
make_table(
    ['Dependency type', 'Detail', 'Constraint / notes'],
    [
        ['AWS integration — Security Hub',
         'Read-only ingestion of findings and standards status (e.g. FSBP, CIS, PCI) to compare live posture against mapped controls.',
         'Non-prod account only. IAM role scoped to securityhub:Get*/Describe*/List* and BatchGet* — read-only.'],
        ['AWS integration — AWS Config',
         'Read-only ingestion of resource configuration items and rule compliance state to evidence control coverage.',
         'Non-prod only. IAM scoped to config:Describe*/Get*/List* and BatchGet*. No config:Put* / remediation.'],
        ['AWS integration — Conformance Packs',
         'Read-only retrieval of conformance-pack compliance summaries mapped to framework controls.',
         'Non-prod only. Read via config:Describe/Get ConformancePack* APIs; packs are not created or modified.'],
        ['Access pattern',
         'Cross-account IAM role assumed by the ComplianceIQ non-prod account, or an exported findings snapshot (S3) shared read-only.',
         'Prefer time-boxed AssumeRole with an external ID; snapshot export is the lower-touch alternative.'],
        ['Client support',
         'Provision the read-only IAM role in the non-prod account (or produce a findings/Config snapshot); confirm in-scope frameworks and account IDs.',
         'A short, well-defined ask; client keeps full control and can revoke at any time.'],
        ['Artefacts needed',
         'Client policy/standard documents for redlining; target framework selection; account/OU scope list; existing control matrix if any.',
         'Accepted as documents or structured export; sensitive data handled per policy.'],
        ['Other assistance',
         'A named client technical contact for the non-prod account; a short scoping session to agree boundaries and success criteria.',
         'Non-invasive; no production access, no write access, no agent installation.'],
    ],
    col_widths=[1.5, 3.5, 1.8],
)
body('AuditBoard (and comparable GRC platforms): if the client uses AuditBoard and a read-only API/'
     'export becomes available, ComplianceIQ could ingest control and evidence records to align its '
     'mapping with the client\'s existing GRC system of record. This is read-only and optional; it is '
     'listed as a future/if-available dependency and is not required for Horizon 2.',
     italic=True, size=9.5, color=GREY, space_after=8)

# ---- 7.3 Horizon 3 dependencies ----
heading('7.3 Horizon 3 Dependencies (Future / Strategic)', level=2, color=CYAN)
body('Horizon 3 broadens coverage and, if pursued, exposes elements to clients. Dependencies grow '
     'accordingly but the non-prod, read-only, non-invasive principles still hold for any environment '
     'integration.', size=10, color=GREY)
make_table(
    ['Dependency type', 'Detail', 'Constraint / notes'],
    [
        ['Multi-region / framework data',
         'Curated regulatory datasets and control crosswalks for new regions (EU, APAC, Americas) and frameworks (PCI DSS, SOC 2, HIPAA, FedRAMP, DORA).',
         'Data-authoring effort; SME validation per new region/framework before use.'],
        ['AWS integration — continuous posture',
         'Scheduled read-only pulls from Security Hub / AWS Config across non-prod accounts for ongoing drift detection.',
         'Same read-only, non-prod IAM posture as H2, extended to a recurring (still read-only) schedule.'],
        ['GRC integration — AuditBoard et al.',
         'Read-only, bi-directional-aware sync (read the client GRC state; never write) to keep evidence aligned.',
         'Only if a supported read-only API exists; write-back explicitly out of scope.'],
        ['Client support (service offering)',
         'Commercial, legal and data-processing agreements; per-tenant isolation sign-off; defined client responsibilities.',
         'Client remains solely responsible for their compliance; tool is advisory, not attestation.'],
        ['Artefacts needed',
         'Per-client scope baselines, tenancy configuration, branding, and data-residency requirements.',
         'Captured during onboarding; residency honoured per client region.'],
        ['Other assistance',
         'Legal/GTM review for client exposure; internal platform security review; SLA and support-model definition.',
         'Governance gates before any client-facing launch.'],
    ],
    col_widths=[1.6, 3.4, 1.8],
)

# ---- 7.4 Integration architecture note ----
heading('7.4 How the Read-Only AWS Integration Works (Reference)', level=2, color=NAVY)
body('For the AWS-service integrations in Horizons 2–3, the recommended pattern is:', space_after=3)
for lead, rest in [
    ('1. Cross-account IAM role — ',
     'the client (or SAS) creates an IAM role in the non-production account that trusts the ComplianceIQ '
     'non-prod account and requires an external ID. The role carries a read-only policy (Get/Describe/'
     'List/BatchGet on Security Hub, Config and Conformance Packs only).'),
    ('2. AssumeRole, time-boxed — ',
     'ComplianceIQ assumes the role for the duration of an assessment via STS, retrieving findings and '
     'configuration state. Credentials are short-lived and can be revoked instantly by the client.'),
    ('3. Read, map, present — ',
     'the retrieved posture is compared against the mapped control set and surfaced to the consultant. '
     'Nothing is written back; there is no remediation action taken in the account.'),
    ('4. Lower-touch alternative — ',
     'where even read-only role assumption is not permitted, the client exports a Security Hub / Config '
     'findings snapshot (e.g. to an S3 location shared read-only) that ComplianceIQ ingests. This '
     'removes any live connection entirely.'),
]:
    bullet(rest, bold_lead=lead)
body('Every option above is non-production, read-only, and non-invasive: ComplianceIQ observes and '
     'informs; the client and the SAS engineer retain sole control of any change to the environment.',
     bold=True, color=NAVY, space_after=8)

# ===========================================================================
# 8. RISKS / CONSIDERATIONS
# ===========================================================================
heading('8. Considerations & Guardrails', level=1)
for lead, rest in [
    ('Advisory, not attestation — ', 'the tool supports consultants; it does not certify compliance. Client responsibility must be explicit in any client-facing use.'),
    ('Data handling — ', 'keep inference inside AWS (already the case with Bedrock); formalise data-residency and retention before any client exposure.'),
    ('Source accuracy — ', 'AI output must remain grounded and citation-backed; retain the human-in-the-loop review that already exists.'),
    ('Internal validation — ', 'validate the Section 4 mapping against the current SAS Target Operating Model (TOM) and go-to-market guidance before external circulation.'),
    ('Security review — ', 'complete an internal application security and privacy review prior to any production or client rollout.'),
]:
    bullet(rest, bold_lead=lead)

# ===========================================================================
# 9. CLOSING
# ===========================================================================
heading('9. Recommendation', level=1)
body('Present ComplianceIQ to the team as a productivity accelerator that already maps cleanly onto '
     'every SAS pillar today, with a credible path to becoming an engagement companion (Horizon 2) and, '
     'in time, a regional and client-facing service (Horizon 3). It reinforces the core SAS value '
     'proposition — turning compliance into a faster, lower-cost business enabler — while remaining '
     'policy-safe inside the AWS boundary.')
body('Suggested next steps: (1) validate the lifecycle mapping against the current TOM with an SSO '
     'holder; (2) prioritise the Horizon 2 backlog with the team; (3) run an internal security/privacy '
     'review; (4) pilot on one live engagement to quantify the time savings referenced in Section 5.',
     bold=True, color=NAVY)

# footer disclaimer
spacer(10)
disc = doc.add_paragraph()
r = disc.add_run('Internal working document. References to AWS SAS content are paraphrased from the '
                 'public AWS Security Assurance Services page and the five deliverables provided; '
                 'internal (SSO-gated) sources were not accessible to the generator and should be '
                 'verified by the team.')
r.font.size = Pt(8.5)
r.font.italic = True
r.font.color.rgb = GREY

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
out_dir = os.path.expanduser('~/Documents')
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, 'ComplianceIQ_SAS_Alignment_and_Roadmap.docx')
doc.save(out_path)
print(f'SAVED: {out_path}')
